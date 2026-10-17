"""
md_to_docx.py -- Render a Markdown manuscript to a submission-style DOCX.

Handles what the manuscript actually uses: headings, paragraphs, bullet / numbered lists
(two levels), block quotes, pipe tables, figures (![..](path)), inline **bold**, *italic*,
`code`, [text](url), Pandoc citations [@key; @key] rendered Harvard author-year from
references.bib, and LaTeX math ($..$ inline, $$..$$ display) rendered as Unicode text with
real Word superscripts / subscripts. Unknown LaTeX commands are reported, not dropped.
"""

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, Pt, RGBColor

BODY_FONT = "Calibri"
MATH_FONT = "Cambria Math"

# --------------------------------------------------------------------------
# Citations
# --------------------------------------------------------------------------

ACCENTS = {"'": "́", "`": "̀", "^": "̂", '"': "̈", "~": "̃", "c": "̧", "v": "̌"}


def bib_accents(name):
    """Turn BibTeX accents such as Ortu\\~no or {\\'e} into composed Unicode characters."""
    import unicodedata
    name = re.sub(r"\\i(?![A-Za-z])", "i", name)  # dotless i, as in Echeverr\'\ia
    out = re.sub(r"\{?\\([\'`^\"~cv])\s*\{?([A-Za-z])\}?\}?",
                 lambda m: unicodedata.normalize("NFC", m.group(2) + ACCENTS[m.group(1)]), name)
    return out


def load_bib(bib_path):
    """key -> (author last names, year) from a BibTeX file."""
    text = Path(bib_path).read_text(encoding="utf-8")
    out = {}
    for m in re.finditer(r"@\w+\{([^,]+),(.*?)\n\}", text, flags=re.S):
        key, body = m.group(1).strip(), m.group(2)
        a = re.search(r"author\s*=\s*\{(.*?)\},?\s*\n", body, flags=re.S)
        y = re.search(r"year\s*=\s*\{?(\d{4})", body)
        names = []
        if a:
            for person in re.split(r"\s+and\s+", a.group(1).replace("\n", " ")):
                person = person.strip().strip("{}")
                last = person.split(",")[0] if "," in person else person.split()[-1]
                names.append(bib_accents(last).replace("{", "").replace("}", "").strip())
        out[key] = (names, y.group(1) if y else "n.d.")
    return out


def cite_text(keys, bib, missing):
    parts = []
    for k in keys:
        if k not in bib:
            missing.add(k)
            parts.append(k)
            continue
        names, year = bib[k]
        if not names:
            who = k
        elif len(names) == 1:
            who = names[0]
        elif len(names) == 2:
            who = f"{names[0]} and {names[1]}"
        else:
            who = f"{names[0]} et al."
        parts.append(f"{who}, {year}")
    return "(" + "; ".join(parts) + ")"

# --------------------------------------------------------------------------
# LaTeX -> runs
# --------------------------------------------------------------------------

SYMBOLS = {
    "alpha": "α", "beta": "β", "gamma": "γ", "delta": "δ", "epsilon": "ε", "varepsilon": "ε",
    "theta": "θ", "lambda": "λ", "mu": "μ", "pi": "π", "rho": "ρ", "sigma": "σ", "tau": "τ",
    "phi": "φ", "varphi": "φ", "omega": "ω", "Delta": "Δ", "Omega": "Ω", "Sigma": "Σ", "Phi": "Φ",
    "times": "×", "cdot": "·", "le": "≤", "leq": "≤", "ge": "≥", "geq": "≥", "ne": "≠", "neq": "≠",
    "approx": "≈", "sim": "~", "pm": "±", "infty": "∞", "in": "∈", "to": "→", "rightarrow": "→",
    "leftarrow": "←", "int": "∫", "sum": "∑", "prod": "∏", "partial": "∂", "mid": "|", "ldots": "…",
    "dots": "…", "cdots": "⋯", "propto": "∝", "ll": "≪", "gg": "≫", "%": "%", "{": "{", "}": "}",
    ",": " ", ";": " ", ":": " ", "!": "", "quad": " ", "qquad": "  ",
    " ": " ", "_": "_", "&": "&", "#": "#", "$": "$", "lvert": "|", "rvert": "|", "langle": "⟨",
    "rangle": "⟩", "left": "", "right": "", "big": "", "Big": "", "bigl": "", "bigr": "",
}
FUNCS = {"cos", "sin", "tan", "arctan", "log", "ln", "exp", "max", "min", "arg", "det", "lim", "sup", "inf"}
TEXT_CMDS = {"text", "mathrm", "operatorname", "textrm", "mathbf", "textbf", "mathit", "textit", "mathsf"}
CALLIGRAPHIC = {"H": "ℋ", "F": "ℱ", "L": "ℒ", "O": "𝒪"}


class LatexRenderer:
    def __init__(self, unknown):
        self.unknown = unknown

    def render(self, src):
        """Return list of (text, vert) with vert in {None, 'sup', 'sub'}."""
        self.s, self.i = src, 0
        return self._seq(None, stop=None)

    def _group(self):
        """Read a {...} group or a single token; return its raw source."""
        s = self.s
        while self.i < len(s) and s[self.i] == " ":
            self.i += 1
        if self.i >= len(s):
            return ""
        if s[self.i] == "{":
            depth, j = 0, self.i
            while j < len(s):
                if s[j] == "{" and (j == 0 or s[j - 1] != "\\"):
                    depth += 1
                elif s[j] == "}" and s[j - 1] != "\\":
                    depth -= 1
                    if depth == 0:
                        break
                j += 1
            inner = s[self.i + 1:j]
            self.i = j + 1
            return inner
        if s[self.i] == "\\":
            m = re.match(r"\\([A-Za-z]+|.)", s[self.i:])
            self.i += m.end()
            return m.group(0)
        ch = s[self.i]
        self.i += 1
        return ch

    def _sub(self, src, vert):
        saved = (self.s, self.i)
        out = self._seq(vert, stop=None, src=src)
        self.s, self.i = saved
        return out

    def _seq(self, vert, stop, src=None):
        if src is not None:
            self.s, self.i = src, 0
        s, out = self.s, []
        while self.i < len(s):
            c = s[self.i]
            if c in "^_":
                self.i += 1
                g = self._group()
                out += self._sub(g, "sup" if c == "^" else "sub")
            elif c == "\\":
                m = re.match(r"\\([A-Za-z]+|.)", s[self.i:])
                self.i += m.end()
                name = m.group(1)
                if name in TEXT_CMDS:
                    out.append((self._group(), vert))
                elif name == "mathcal":
                    g = self._group()
                    out.append(("".join(CALLIGRAPHIC.get(ch, ch) for ch in g), vert))
                elif name in ("frac", "dfrac", "tfrac"):
                    num, den = self._group(), self._group()
                    wrap = lambda x: x if re.fullmatch(r"[\w.]+", x) else "(" + x + ")"
                    out += self._sub(wrap(num), vert) + [("/", vert)] + self._sub(wrap(den), vert)
                elif name == "sqrt":
                    out += [("√(", vert)] + self._sub(self._group(), vert) + [(")", vert)]
                elif name in FUNCS:
                    out.append((name, vert))
                elif name in SYMBOLS:
                    out.append((SYMBOLS[name], vert))
                else:
                    self.unknown.add(name)
                    out.append((name, vert))
            elif c in "{}":
                self.i += 1
            elif c == "~":
                out.append((" ", vert))
                self.i += 1
            else:
                j = self.i
                while j < len(s) and s[j] not in "^_\\{}~":
                    j += 1
                out.append((s[self.i:j], vert))
                self.i = j
        return out

# --------------------------------------------------------------------------
# Inline markdown -> runs
# --------------------------------------------------------------------------

INLINE = re.compile(
    r"(?P<math>(?<!\\)\$(?!\$)(?P<m>.+?)(?<!\\)\$)"
    r"|(?P<cite>\[(?P<c>@[^\]]+)\])"
    r"|(?P<link>(?<!!)\[(?P<lt>[^\]]+)\]\((?P<lu>[^)]+)\))"
    r"|(?P<code>`(?P<k>[^`]+)`)"
    r"|(?P<bold>\*\*(?P<b>.+?)\*\*)"
    r"|(?P<ital>(?<![\w*])\*(?P<i>[^*\s][^*]*?)\*(?![\w*]))"
)


class Inline:
    def __init__(self, bib):
        self.bib = bib
        self.missing = set()
        self.unknown = set()
        self.latex = LatexRenderer(self.unknown)

    def runs(self, text, bold=False, italic=False):
        """Yield dicts: text, bold, italic, code, math, vert."""
        pos, out = 0, []
        for m in INLINE.finditer(text):
            if m.start() > pos:
                out.append(dict(text=text[pos:m.start()], bold=bold, italic=italic))
            if m.group("math"):
                for t, v in self.latex.render(m.group("m")):
                    out.append(dict(text=t, bold=bold, italic=italic, math=True, vert=v))
            elif m.group("cite"):
                keys = re.findall(r"@([\w\-:]+)", m.group("c"))
                out.append(dict(text=cite_text(keys, self.bib, self.missing), bold=bold, italic=italic))
            elif m.group("link"):
                out += self.runs(m.group("lt"), bold, italic)
            elif m.group("code"):
                out.append(dict(text=m.group("k"), bold=bold, italic=italic, code=True))
            elif m.group("bold"):
                out += self.runs(m.group("b"), True, italic)
            elif m.group("ital"):
                out += self.runs(m.group("i"), bold, True)
            pos = m.end()
        if pos < len(text):
            out.append(dict(text=text[pos:], bold=bold, italic=italic))
        for r in out:
            r["text"] = r["text"].replace("\\%", "%").replace("\\$", "$")
        return out

    def add(self, paragraph, text, size=11, bold=False, italic=False, color=None):
        for r in self.runs(text, bold, italic):
            if not r["text"]:
                continue
            run = paragraph.add_run(r["text"])
            run.font.size = Pt(size)
            run.font.bold = r.get("bold") or None
            run.font.italic = r.get("italic") or None
            run.font.name = MATH_FONT if r.get("math") else ("Consolas" if r.get("code") else BODY_FONT)
            if r.get("vert") == "sup":
                run.font.superscript = True
            elif r.get("vert") == "sub":
                run.font.subscript = True
            if color:
                run.font.color.rgb = color
            if r.get("code") and r["text"].startswith("["):
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW  # author placeholder: must be filled in
        return paragraph

# --------------------------------------------------------------------------
# Blocks
# --------------------------------------------------------------------------

def build(md_path, docx_path, bib_path):
    md_path = Path(md_path)
    inl = Inline(load_bib(bib_path))
    doc = Document()
    for sec in doc.sections:
        sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1.0)
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(11)

    lines = md_path.read_text(encoding="utf-8").splitlines()
    i = 0

    def para(space_after=6):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        return p

    def heading(text, level):
        h = doc.add_heading(level=level)
        inl.add(h, text, size={0: 18, 1: 15, 2: 13, 3: 12}.get(level, 11), color=RGBColor(0x1F, 0x3A, 0x5F))
        return h

    while i < len(lines):
        raw = lines[i]
        line = raw.strip()

        if not line or re.fullmatch(r"-{3,}|\*{3,}", line):
            i += 1
            continue

        # Display math, possibly spanning lines
        if line.startswith("$$"):
            buf = line[2:]
            while not buf.rstrip().endswith("$$") and i + 1 < len(lines):
                i += 1
                buf += " " + lines[i].strip()
            p = para(8)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            inl.add(p, "$" + buf.rstrip()[:-2].strip() + "$", size=11.5)
            i += 1
            continue

        # Figure
        m = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", line)
        if m:
            img = (md_path.parent / m.group(2)).resolve()
            if img.exists():
                doc.add_picture(str(img), width=Inches(6.3))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                inl.add(para(), f"[Missing figure file: {m.group(2)}]", italic=True)
            i += 1
            continue

        # Table
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            ncol = max(len(r) for r in rows)
            tbl = doc.add_table(rows=len(rows), cols=ncol)
            tbl.style = "Table Grid"
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            for r, cells in enumerate(rows):
                for c in range(ncol):
                    cell = tbl.cell(r, c)
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(2)
                    inl.add(p, cells[c] if c < len(cells) else "", size=9, bold=(r == 0))
                    if r == 0:
                        cell._tc.get_or_add_tcPr().append(parse_xml(r'<w:shd {} w:fill="EBF2F7"/>'.format(nsdecls("w"))))
            para(4)
            continue

        # Headings
        m = re.match(r"(#{1,4})\s+(.*)", line)
        if m:
            heading(m.group(2), len(m.group(1)) - 1)
            i += 1
            continue

        # Lists (indent >= 2 spaces -> level 2)
        m = re.match(r"(\s*)([-*]|\d+\.)\s+(.*)", raw)
        if m:
            indent, marker, body = m.groups()
            level2 = len(indent.replace("\t", "    ")) >= 2
            if marker[0].isdigit():
                # Literal numbers: Word's List Number style would keep counting across the document
                p = para(2)
                p.paragraph_format.left_indent = Inches(0.6 if level2 else 0.35)
                p.paragraph_format.first_line_indent = Inches(-0.25)
                inl.add(p, marker + " " + body)
            else:
                p = doc.add_paragraph(style="List Bullet" + (" 2" if level2 else ""))
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(2)
                inl.add(p, body)
            i += 1
            continue

        # Block quote
        if line.startswith(">"):
            p = para(4)
            p.paragraph_format.left_indent = Inches(0.4)
            inl.add(p, line.lstrip("> "), size=10.5, italic=True)
            i += 1
            continue

        # Figure caption
        if re.match(r"\*\*(Figure|Table) \d+\.", line):
            p = para(10)
            inl.add(p, line, size=9.5, color=RGBColor(0x33, 0x44, 0x66))
            i += 1
            continue

        # Paragraph: join soft-wrapped lines; a line ending in two spaces is a hard break
        p = para()
        inl.add(p, line)
        hard = raw.endswith("  ")
        i += 1
        while i < len(lines) and lines[i].strip() and not hard and not re.match(
                r"\s*(#|\||!\[|\$\$|>|[-*]\s|\d+\.\s|---)", lines[i]):
            p.add_run(" ")
            inl.add(p, lines[i].strip())
            hard = lines[i].endswith("  ")
            i += 1

    doc.save(str(docx_path))
    return {"missing_citations": sorted(inl.missing), "unknown_latex": sorted(inl.unknown)}


if __name__ == "__main__":
    import argparse
    import sys
    ap = argparse.ArgumentParser(description="Render a Markdown manuscript to DOCX.")
    ap.add_argument("markdown")
    ap.add_argument("docx")
    ap.add_argument("bib", help="references.bib used for author-year citations")
    a = ap.parse_args()
    res = build(a.markdown, a.docx, a.bib)
    print(res)
    # Unresolved citations or unknown LaTeX would show up as raw text in the DOCX
    sys.exit(1 if res["missing_citations"] or res["unknown_latex"] else 0)
