"""
audit_deliverables.py -- Text checks that caught real defects in the HHSA run.

Checks, over HTML / JS / JSON / MD / BIB / DOCX / PPTX files under the given paths:
  1. privacy   : forbidden strings (confidential manuscript IDs, raw file stems, calendar dates,
                 participant codes) -- pass them with --forbid / --forbid-regex
  2. stale     : numbers or claims known to be wrong (old fabricated values) -- --stale
  3. markup    : leftover Markdown / LaTeX / Pandoc citations inside DOCX and PPTX text
  4. links     : local href/src targets in HTML that do not exist (relative to the file)
  5. css       : class names used in HTML but defined in no stylesheet or <style> block

Usage:
  python audit_deliverables.py site blog deck paper/review.docx \
      --forbid <CONFIDENTIAL-MS-ID> <raw_file_stem> <YYYY-MM-DD> --forbid-regex "\\bS1\\d\\d\\b" \
      --stale "714 beats" "100% recovery"
Run the link check on the staged repo clone (repo layout), or pass --link-ignore for paths
that only exist there. Classes styled by external CSS or used only by JS go in --class-ignore.
Exit code 1 if anything is found.
"""

import argparse
import re
import sys
from pathlib import Path

TEXT_EXT = {".html", ".htm", ".js", ".css", ".json", ".md", ".bib", ".txt", ".csv"}
MARKUP = [r"\*\*[^*\s]", r"(?<!\\)\$[^$\s][^$]*\$", r"\[@\w", r"\\text\{", r"\\frac\{", r"!\[[^\]]*\]\(", r"^#{1,4} "]


def docx_text(p):
    import docx
    d = docx.Document(str(p))
    parts = [para.text for para in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts += [c.text for c in row.cells]
    return "\n".join(parts)


def pptx_text(p):
    from pptx import Presentation
    prs = Presentation(str(p))
    parts = []
    for s in prs.slides:
        for sh in s.shapes:
            if sh.has_text_frame:
                parts.append(sh.text_frame.text)
        if s.has_notes_slide:
            parts.append(s.notes_slide.notes_text_frame.text)
    return "\n".join(parts)


def files(paths):
    for root in paths:
        root = Path(root)
        items = [root] if root.is_file() else root.rglob("*")
        for f in items:
            if f.is_file() and f.name != "desktop.ini" and (f.suffix.lower() in TEXT_EXT | {".docx", ".pptx"}):
                yield f


def read(f):
    if f.suffix.lower() == ".docx":
        return docx_text(f)
    if f.suffix.lower() == ".pptx":
        return pptx_text(f)
    return f.read_text(encoding="utf-8", errors="ignore")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("paths", nargs="+")
    ap.add_argument("--forbid", nargs="*", default=[])
    ap.add_argument("--forbid-regex", nargs="*", default=[])
    ap.add_argument("--stale", nargs="*", default=[])
    ap.add_argument("--link-ignore", nargs="*", default=[],
                    help="substrings of link targets to skip (e.g. ../slides/ when auditing the working layout)")
    ap.add_argument("--class-ignore", nargs="*", default=[],
                    help="class names used only as JS hooks or styled by external CSS")
    a = ap.parse_args()

    problems = []
    css_defined, html_classes = set(), {}
    for f in files(a.paths):
        text = read(f)
        for t in a.forbid:
            if t in text:
                problems.append(f"privacy  {f}: '{t}'")
        for rx in a.forbid_regex:
            for m in set(re.findall(rx, text)):
                problems.append(f"privacy  {f}: /{rx}/ -> {m}")
        for t in a.stale:
            if t in text:
                problems.append(f"stale    {f}: '{t}'")
        if f.suffix.lower() in (".docx", ".pptx"):
            for rx in MARKUP:
                m = re.search(rx, text, flags=re.M)
                if m:
                    problems.append(f"markup   {f}: raw markup {m.group(0)[:30]!r}")
        if f.suffix.lower() == ".css":
            css_defined |= set(re.findall(r"\.([A-Za-z_][\w-]*)", text))
        if f.suffix.lower() in (".html", ".htm"):
            for style in re.findall(r"<style[^>]*>(.*?)</style>", text, flags=re.S):
                css_defined |= set(re.findall(r"\.([A-Za-z_][\w-]*)", style))
            for cl in re.findall(r'class="([^"$`{]*)"', text):
                for c in cl.split():
                    html_classes.setdefault(c, f)
            for tgt in re.findall(r'(?:href|src)="([^"#?]+)', text):
                if re.match(r"(https?:|data:|mailto:|javascript:|//)", tgt) or "${" in tgt                         or any(i in tgt for i in a.link_ignore):
                    continue
                if not (f.parent / tgt).exists():
                    problems.append(f"link     {f}: missing local target {tgt}")
    for c, f in sorted(html_classes.items()):
        # Font Awesome (fa, fas, fa-*) is styled by its CDN stylesheet, which is not scanned
        if c not in css_defined and c not in a.class_ignore and not re.fullmatch(r"fa[srbl]?|fa-[\w-]+", c):
            problems.append(f"css      {f}: class '{c}' has no CSS rule")

    print("\n".join(problems) if problems else "audit clean")
    sys.exit(1 if problems else 0)


if __name__ == "__main__":
    main()
