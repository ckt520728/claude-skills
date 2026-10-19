"""MediCloud screenshot batch tool.

Takes a folder of screenshots, trims empty borders (and optional browser chrome),
and places each one full-width on a landscape page of the MediCloud Word template.

SAFE BY DESIGN: original screenshots are never modified or moved. Cropped copies
go to <folder>\\_cropped and the Word file is written to a NEW path (it refuses to
overwrite an existing file unless --force is given).

Usage:
  python medicloud_batch.py FOLDER --name 王小明 --chart 1234567 [options]
Options:
  --template PATH   template .docx (default: MediCloud-健保雲端摘要範本.docx next to this script)
  --out PATH        output .docx (default: <FOLDER>\\<YYYY-MM-DD>-<name>.docx)
  --date YYYY-MM-DD query date (default: today)
  --by NAME         查詢者
  --top N           strip N px from the top of every image first (browser toolbar)
  --bottom N        strip N px from the bottom of every image first
  --left N          strip N px from the left edge (e.g. browser window frame)
  --right N         strip N px from the right edge
  --no-trim         skip automatic border trimming
  --scale X         upscale cropped copies by X (e.g. 2) for smoother Word rendering
  --dry-run         show what would happen, write nothing
  --force           allow overwriting --out
"""
import argparse
import datetime as dt
import os
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageChops

EXTS = {".png", ".jpg", ".jpeg", ".bmp"}
PAGE_W_CM, PAGE_H_CM = 27.1, 16.2  # usable landscape area under heading + caption
GREY = RGBColor(0x66, 0x66, 0x66)


def natural_key(p: Path):
    return [int(t) if t.isdigit() else t.lower() for t in re.split(r"(\d+)", p.name)]


def trim(img: Image.Image, tol=12, pad=14):
    """Crop away a uniform border (colour taken from the top-left pixel)."""
    rgb = img.convert("RGB")
    bg = Image.new("RGB", rgb.size, rgb.getpixel((0, 0)))
    diff = ImageChops.difference(rgb, bg).convert("L").point(lambda v: 255 if v > tol else 0)
    box = diff.getbbox()
    if not box:
        return img
    l, t, r, b = box
    l, t = max(0, l - pad), max(0, t - pad)
    r, b = min(img.width, r + pad), min(img.height, b + pad)
    return img.crop((l, t, r, b))


def fit_cm(w_px, h_px):
    """Width/height in cm that fit the page box, preserving aspect ratio."""
    w = PAGE_W_CM
    h = w * h_px / w_px
    if h > PAGE_H_CM:
        h = PAGE_H_CM
        w = h * w_px / h_px
    return w, h


def set_header_text(sec, name, chart, date_str):
    hp = sec.header.paragraphs[0]
    for r in hp.runs[1:]:
        r._r.getparent().remove(r._r)
    if hp.runs:
        hp.runs[0].text = f"病人：{name or '__________'}  病歷號：{chart or '__________'}  查詢日期：{date_str}"


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("folder")
    ap.add_argument("--name", default="")
    ap.add_argument("--chart", default="")
    ap.add_argument("--by", default="")
    ap.add_argument("--date", default=dt.date.today().isoformat())
    here = Path(__file__).resolve().parent
    tpl = next((c for c in (here / "MediCloud-健保雲端摘要範本.docx",
                            here.parent / "assets" / "MediCloud-健保雲端摘要範本.docx") if c.exists()),
               here / "MediCloud-健保雲端摘要範本.docx")
    ap.add_argument("--template", default=str(tpl))
    ap.add_argument("--out", default="")
    ap.add_argument("--top", type=int, default=0)
    ap.add_argument("--bottom", type=int, default=0)
    ap.add_argument("--left", type=int, default=0)
    ap.add_argument("--right", type=int, default=0)
    ap.add_argument("--no-trim", action="store_true")
    ap.add_argument("--scale", type=float, default=1.0)
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--force", action="store_true")
    a = ap.parse_args()

    folder = Path(a.folder)
    if not folder.is_dir():
        sys.exit(f"Folder not found: {folder}")
    files = sorted((p for p in folder.iterdir() if p.is_file() and p.suffix.lower() in EXTS), key=natural_key)
    if not files:
        sys.exit("No images (.png/.jpg/.bmp) found in that folder.")

    safe_name = re.sub(r'[\\/:*?"<>|]', "", a.name) or "patient"
    out = Path(a.out) if a.out else folder / f"{a.date}-{safe_name}.docx"
    if out.exists() and not a.force:
        sys.exit(f"Refusing to overwrite existing file: {out}\n(use --force, or pass a different --out)")
    crop_dir = folder / "_cropped"

    print(f"{len(files)} image(s) -> {out}")
    if a.dry_run:
        for f in files:
            print("  would process:", f.name)
        print("Dry run: nothing written.")
        return

    crop_dir.mkdir(exist_ok=True)
    doc = Document(a.template)

    # fill cover info
    set_header_text(doc.sections[0], a.name, a.chart, a.date)
    set_header_text(doc.sections[1], a.name, a.chart, a.date)
    vals = {"病人姓名": a.name, "病歷號": a.chart, "查詢日期": a.date, "查詢者": a.by}
    tbl = doc.tables[0]
    for row in tbl.rows:
        for i in (0, 2):
            key = row.cells[i].text.strip()
            if key in vals and vals[key]:
                row.cells[i + 1].text = vals[key]

    # drop the template's placeholder pages (everything after the section-1 break paragraph)
    body = doc.element.body
    children = list(body)
    brk = next(i for i, el in enumerate(children)
               if el.tag.endswith("}p") and el.find(".//{*}sectPr") is not None)
    for el in children[brk + 1:]:
        if not el.tag.endswith("}sectPr"):
            body.remove(el)

    for n, f in enumerate(files):
        img = Image.open(f)
        if a.top or a.bottom or a.left or a.right:
            box = (a.left, a.top, img.width - a.right, img.height - a.bottom)
            if box[2] <= box[0] or box[3] <= box[1]:
                sys.exit(f"{f.name}: --top/--bottom/--left/--right strip more than the image size {img.size}")
            img = img.crop(box)
        if not a.no_trim:
            img = trim(img)
        if a.scale and a.scale != 1.0:
            img = img.resize((int(img.width * a.scale), int(img.height * a.scale)), Image.LANCZOS)
        cropped = crop_dir / (f.stem + ".png")
        img.convert("RGB").save(cropped, dpi=(192, 192))

        title = re.sub(r"^\d+[\s_.\-]*", "", f.stem) or f.stem
        h = doc.add_heading(title, level=1)
        if n:
            h.paragraph_format.page_break_before = True
        w_cm, h_cm = fit_cm(img.width, img.height)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(cropped), width=Cm(w_cm), height=Cm(h_cm))
        taken = dt.datetime.fromtimestamp(f.stat().st_mtime).strftime("%Y-%m-%d %H:%M")
        cap = doc.add_paragraph(f"來源檔：{f.name}　擷取時間：{taken}", style="Caption")
        print(f"  {f.name}: {img.width}x{img.height}px -> {w_cm:.1f}x{h_cm:.1f} cm "
              f"({img.width / (w_cm / 2.54):.0f} px/inch)")

    doc.save(out)
    print("Saved:", out)
    print("Originals untouched; cropped copies in:", crop_dir)


if __name__ == "__main__":
    main()
