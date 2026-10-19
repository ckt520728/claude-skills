"""Reorganize a pasted MediCloud 西醫用藥紀錄 Word file into a table-free summary.

Reads the medication table in SOURCE (read-only), writes a NEW .docx with:
  - patient banner (as shown on the page)
  - Part A: by visit date -> diagnosis -> drug bullets
  - Part B: by drug -> dispensing dates
Identical rows are kept and marked with their original 項次 numbers (nothing dropped).
Refuses to overwrite an existing output file.

Usage: python medicloud_med_reorganize.py SOURCE.docx [OUT.docx]
"""
import re
import sys
from collections import OrderedDict
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "Microsoft JhengHei"
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
GREY = RGBColor(0x66, 0x66, 0x66)
COLS = ["項次", "就醫日期", "來源", "主診斷", "ATC3名稱", "複方註記", "成分名稱",
        "藥品名稱", "用法用量", "藥品用量", "給藥日數", "單筆餘藥日數試算"]


def uniq_cells(row):
    seen, out = set(), []
    for c in row.cells:
        if id(c._tc) in seen:
            continue
        seen.add(id(c._tc))
        out.append(c.text.strip())
    return out


def roc_to_iso(s):
    m = re.fullmatch(r"(\d{2,3})/(\d{2})/(\d{2})", s)
    return f"{int(m.group(1)) + 1911}-{m.group(2)}-{m.group(3)}" if m else s


def parse(src):
    d = Document(src)
    if not d.tables:
        sys.exit("No table found in the source file.")
    t = d.tables[0]
    hdr = uniq_cells(t.rows[0])
    if hdr != COLS:
        sys.exit(f"Unexpected table header:\n  {hdr}\nexpected:\n  {COLS}")
    recs = []
    for r in t.rows[1:]:
        c = uniq_cells(r)
        if not any(c):
            continue
        if len(c) != len(COLS):
            sys.exit(f"Row has {len(c)} cells, expected {len(COLS)}: {c}")
        rec = dict(zip(COLS, c))
        src_parts = [p.strip() for p in rec["來源"].split("\n") if p.strip()] or [rec["來源"]]
        rec["院所"] = " ".join(src_parts)
        dx = [p.strip() for p in rec["主診斷"].split("\n") if p.strip()]
        rec["診斷"] = dx[0] if dx else ""
        rec["ICD"] = dx[1] if len(dx) > 1 else ""
        recs.append(rec)
    # patient banner = first paragraph that carries the ID line
    banner = next((p.text.strip() for p in d.paragraphs if "身分證號" in p.text), "")
    banner = banner.split("查詢其他")[0].strip()  # drop web-page button text
    note = next((p.text.strip() for p in d.paragraphs if "家醫計畫" in p.text), "")
    return recs, banner, note


def set_font(style, size, bold=None, color=None):
    style.font.name = FONT
    style.font.size = Pt(size)
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        rf.set(qn(a), FONT)
    if bold is not None:
        style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color


def drug_line(r):
    ing = r["成分名稱"] or r["藥品名稱"]
    tag = "（複方）" if r["複方註記"].upper() == "Y" else ""
    return f"{ing}{tag}", r["藥品名稱"]


def detail(r):
    return (f"用法 {r['用法用量']}　用量 {r['藥品用量']}　給藥 {r['給藥日數']} 日　"
            f"餘藥試算 {r['單筆餘藥日數試算']} 日　分類：{r['ATC3名稱']}")


def build(recs, banner, note, out, src_name):
    doc = Document()
    set_font(doc.styles["Normal"], 11)
    set_font(doc.styles["Title"], 20, True, NAVY)
    set_font(doc.styles["Heading 1"], 14, True, NAVY)
    set_font(doc.styles["Heading 2"], 12, True, NAVY)
    set_font(doc.styles["List Bullet"], 11)
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(1.9)
    s.top_margin = s.bottom_margin = Cm(1.8)

    fp = s.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("資料來源：健保醫療資訊雲端查詢系統（NHI MediCloud）· 僅供本院診療使用，請勿外流")
    r.font.size = Pt(9)
    r.font.color.rgb = GREY

    doc.add_paragraph("健保雲端西醫用藥紀錄（整理版）", style="Title")
    if banner:
        p = doc.add_paragraph()
        p.add_run(banner).bold = True
    if note:
        doc.add_paragraph(note)
    dates = list(OrderedDict.fromkeys(r["就醫日期"] for r in recs))
    p = doc.add_paragraph()
    p.add_run(f"共 {len(recs)} 筆用藥、{len(dates)} 個就醫日"
              f"（{roc_to_iso(dates[-1])} ～ {roc_to_iso(dates[0])}）；原始檔：{src_name}").font.color.rgb = GREY

    # ---- Part A: by visit date
    doc.add_heading("A. 依就醫日期", level=1)
    visits = OrderedDict()
    for r in recs:
        visits.setdefault((r["就醫日期"], r["院所"], r["診斷"], r["ICD"]), []).append(r)
    for (date, site, dx, icd), rows in visits.items():
        doc.add_heading(f"{date}（{roc_to_iso(date)}）　{site}", level=2)
        p = doc.add_paragraph()
        p.add_run("診斷：").bold = True
        p.add_run(f"{dx}　{icd}")
        # collapse identical rows within the visit, keep 項次
        groups = OrderedDict()
        for r in rows:
            key = tuple(r[k] for k in COLS[1:])
            groups.setdefault(key, []).append(r)
        for grp in groups.values():
            r0 = grp[0]
            name, brand = drug_line(r0)
            nums = "、".join(g["項次"] for g in grp)
            rep = f"　【同樣紀錄 ×{len(grp)}，項次 {nums}】" if len(grp) > 1 else f"　（項次 {nums}）"
            b = doc.add_paragraph(style="List Bullet")
            b.add_run(name).bold = True
            b.add_run(f" — {brand}")
            b.add_run(rep).font.color.rgb = GREY
            d = doc.add_paragraph(detail(r0))
            d.paragraph_format.left_indent = Cm(1.27)
            d.paragraph_format.space_after = Pt(4)
            for run in d.runs:
                run.font.size = Pt(10)
                run.font.color.rgb = GREY

    # ---- Part B: by drug
    doc.add_heading("B. 依藥品彙整（各藥配發日期）", level=1)
    by_drug = OrderedDict()
    for r in recs:
        by_drug.setdefault((r["成分名稱"] or r["藥品名稱"], r["藥品名稱"], r["用法用量"]), []).append(r)
    for (ing, brand, usage), rows in by_drug.items():
        b = doc.add_paragraph(style="List Bullet")
        b.add_run(ing).bold = True
        tag = "（複方）" if rows[0]["複方註記"].upper() == "Y" else ""
        b.add_run(f"{tag} — {brand}　用法 {usage}")
        per_date = OrderedDict()
        for r in rows:
            per_date.setdefault(r["就醫日期"], []).append(r)
        parts = []
        for date, rs in per_date.items():
            days = "+".join(x["給藥日數"] for x in rs)
            parts.append(f"{date}（{roc_to_iso(date)}）{len(rs)} 筆，給藥日數 {days}")
        d = doc.add_paragraph("；".join(parts))
        d.paragraph_format.left_indent = Cm(1.27)
        d.paragraph_format.space_after = Pt(4)
        for run in d.runs:
            run.font.size = Pt(10)
            run.font.color.rgb = GREY

    doc.core_properties.title = "健保雲端西醫用藥紀錄（整理版）"
    doc.save(out)


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    src = Path(sys.argv[1])
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else src.with_name(src.stem + " 整理版.docx")
    if out.exists():
        sys.exit(f"Refusing to overwrite existing file: {out}")
    recs, banner, note = parse(src)
    build(recs, banner, note, out, src.name)
    print(f"{len(recs)} records -> {out}")


if __name__ == "__main__":
    main()
